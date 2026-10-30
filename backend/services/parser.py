"""
Multi-format document parser.
Supports: PDF (with OCR for scanned pages), DOCX, PPTX, ODT, RTF, CSV, XLSX/XLS, TXT/MD/RST/LOG,
HTML, JSON, images (OCR) and web URLs.
Returns a list of {page_number, text, metadata} dicts ("pages").
"""

from __future__ import annotations

import csv
import json
import logging
import re
import threading
import zipfile
from pathlib import Path
from typing import Any, Optional

from backend.client.web_client import WebPageClient
from backend.core.exceptions import DocumentParseError

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff", "gif"}
DOCUMENT_EXTENSIONS = {"pdf", "docx", "pptx", "odt", "rtf", "csv", "xlsx", "xls", "txt", "md",
                       "rst", "html", "htm", "json", "log"}
SUPPORTED_EXTENSIONS = DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
MAX_PDF_PAGES = 1000
MAX_OCR_PDF_PAGES = 30          # scanned pages OCR'd per PDF (OCR is CPU-heavy)
MAX_IMAGE_SIDE = 2000           # downscale larger images before OCR
MIN_PAGE_CHARS = 30
MIN_OCR_CHARS = 10              # OCR'd pages are short (slides, receipts); keep anything real

_OCR_LOCK = threading.Lock()
_OCR_ENGINE: Any = None
_OCR_FAILED = False

Page = dict[str, Any]


class DocumentParser:
    def __init__(self, web_client: Optional[WebPageClient] = None) -> None:
        self._web = web_client or WebPageClient()

    # ── public API ───────────────────────────────────────────────────────────

    def parse_file(self, file_path: str, filename: str) -> list[Page]:
        ext = Path(filename).suffix.lower().lstrip(".")
        if ext not in SUPPORTED_EXTENSIONS:
            raise DocumentParseError(
                f"Unsupported file type '.{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
        path = Path(file_path)
        handlers = {
            "pdf": self._parse_pdf, "docx": self._parse_docx, "csv": self._parse_csv,
            "xlsx": self._parse_excel, "xls": self._parse_excel, "txt": self._parse_text,
            "md": self._parse_text, "rst": self._parse_text, "log": self._parse_text,
            "pptx": self._parse_pptx,
            "odt": self._parse_odt, "rtf": self._parse_rtf, "html": self._parse_html,
            "htm": self._parse_html, "json": self._parse_json,
        }
        if ext in IMAGE_EXTENSIONS:
            return self._parse_image(path)
        return handlers[ext](path)

    def parse_url(self, url: str) -> list[Page]:
        html = self._web.fetch_html(url)
        text, parser_name = self._extract_article(html)
        if not text or len(text.strip()) < 100:
            raise DocumentParseError("No readable article text found at that URL.")
        return [
            {"page_number": i + 1, "text": chunk,
             "metadata": {"parser": parser_name, "url": url}}
            for i, chunk in enumerate(self._split_into_chunks(text, 1200))
        ]

    @staticmethod
    def _extract_article(html: str) -> tuple[str, str]:
        try:
            import trafilatura

            text = trafilatura.extract(html, include_tables=True, include_links=False)
            if text and len(text.strip()) > 100:
                return text, "trafilatura"
        except Exception as exc:  # noqa: BLE001 — fall through to BeautifulSoup
            logger.debug("trafilatura failed: %s", exc)
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(html, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
                tag.decompose()
            text = re.sub(r"\n{3,}", "\n\n", soup.get_text(separator="\n", strip=True))
            return text, "beautifulsoup"
        except Exception as exc:  # noqa: BLE001
            logger.warning("BeautifulSoup failed: %s", exc)
        return "", "none"

    # ── file parsers ─────────────────────────────────────────────────────────

    def _parse_pdf(self, path: Path) -> list[Page]:
        """Text layer via pdfium (fast); pdfplumber/PyPDF2 only if pdfium cannot read the file."""
        pages: list[Page] = []
        blank: list[int] = []            # pages with (almost) no text layer → OCR candidates
        parsed = False
        for name, extract in (("pdfium", self._pdf_pdfium), ("pdfplumber", self._pdf_pdfplumber),
                              ("PyPDF2", self._pdf_pypdf2)):
            try:
                pages, blank = extract(path)
                parsed = True
                break                    # a readable file with no text goes to OCR, not other parsers
            except Exception as exc:  # noqa: BLE001 — try the next parser
                logger.warning("%s failed: %s", name, exc)
        if not parsed:
            raise DocumentParseError("The PDF could not be read (corrupt or encrypted).")

        if blank:
            pages.extend(self._ocr_pdf_pages(path, blank))
            pages.sort(key=lambda p: p["page_number"])
        return pages

    @staticmethod
    def _pdf_pdfium(path: Path) -> tuple[list[Page], list[int]]:
        import pypdfium2 as pdfium

        pages: list[Page] = []
        blank: list[int] = []
        pdf = pdfium.PdfDocument(str(path))
        try:
            for num in range(1, min(len(pdf), MAX_PDF_PAGES) + 1):
                textpage = pdf[num - 1].get_textpage()
                try:
                    text = textpage.get_text_bounded().strip()
                finally:
                    textpage.close()
                if len(text) >= MIN_PAGE_CHARS:
                    pages.append({"page_number": num, "text": text,
                                  "metadata": {"parser": "pdfium"}})
                else:
                    blank.append(num)
        finally:
            pdf.close()
        return pages, blank

    @staticmethod
    def _pdf_pdfplumber(path: Path) -> tuple[list[Page], list[int]]:
        import pdfplumber

        pages: list[Page] = []
        blank: list[int] = []
        with pdfplumber.open(str(path)) as pdf:
            for num, page in enumerate(pdf.pages[:MAX_PDF_PAGES], start=1):
                text = (page.extract_text() or "").strip()   # includes table text
                if len(text) >= MIN_PAGE_CHARS:
                    pages.append({"page_number": num, "text": text,
                                  "metadata": {"parser": "pdfplumber"}})
                else:
                    blank.append(num)
        return pages, blank

    @staticmethod
    def _pdf_pypdf2(path: Path) -> tuple[list[Page], list[int]]:
        import PyPDF2

        pages: list[Page] = []
        blank: list[int] = []
        with open(path, "rb") as fh:
            reader = PyPDF2.PdfReader(fh)
            for num, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
                text = (page.extract_text() or "").strip()
                if len(text) >= MIN_PAGE_CHARS:
                    pages.append({"page_number": num, "text": text,
                                  "metadata": {"parser": "PyPDF2"}})
                else:
                    blank.append(num)
        return pages, blank

    def _ocr_pdf_pages(self, path: Path, page_numbers: list[int]) -> list[Page]:
        """OCR scanned PDF pages. Best effort: returns [] when OCR is unavailable."""
        engine = _get_ocr()
        if engine is None:
            logger.info("PDF has %d image-only pages but OCR is unavailable", len(page_numbers))
            return []
        out: list[Page] = []
        try:
            import numpy as np
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            try:
                for num in page_numbers[:MAX_OCR_PDF_PAGES]:
                    image = pdf[num - 1].render(scale=1.5).to_pil().convert("RGB")
                    text = _ocr_array(engine, np.asarray(_limit_side(image)))
                    if len(text) >= MIN_OCR_CHARS:
                        out.append({"page_number": num, "text": text,
                                    "metadata": {"parser": "ocr", "ocr": True}})
            finally:
                pdf.close()
        except Exception as exc:  # noqa: BLE001
            logger.warning("PDF OCR failed: %s", exc)
        if len(page_numbers) > MAX_OCR_PDF_PAGES:
            logger.info("OCR limited to first %d of %d scanned pages",
                        MAX_OCR_PDF_PAGES, len(page_numbers))
        return out

    def _parse_image(self, path: Path) -> list[Page]:
        engine = _get_ocr()
        if engine is None:
            raise DocumentParseError(
                "Image text extraction (OCR) is not available on this server.")
        try:
            import numpy as np
            from PIL import Image

            with Image.open(str(path)) as img:
                img.seek(0)
                rgb = _limit_side(img.convert("RGB"))
            text = _ocr_array(engine, np.asarray(rgb))
        except DocumentParseError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The image could not be read.") from exc
        if not text.strip():
            raise DocumentParseError("No readable text was found in the image.")
        return [{"page_number": 1, "text": text,
                 "metadata": {"parser": "ocr", "ocr": True}}]

    def _parse_docx(self, path: Path) -> list[Page]:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentParseError("DOCX support is not installed.") from exc
        try:
            doc = Document(str(path))
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The DOCX file could not be read.") from exc

        pages: list[Page] = []
        current: list[str] = []
        section = 1
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style is not None and para.style.name.startswith("Heading") and current:
                pages.append({"page_number": section, "text": "\n".join(current),
                              "metadata": {"parser": "python-docx", "type": "section"}})
                section += 1
                current = [text]
            else:
                current.append(text)
        for table in doc.tables:
            rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            current.append("\n".join(rows))
        if current:
            pages.append({"page_number": section, "text": "\n".join(current),
                          "metadata": {"parser": "python-docx", "type": "section"}})
        return pages

    def _parse_csv(self, path: Path) -> list[Page]:
        try:
            with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
                rows = list(csv.reader(fh))
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The CSV file could not be read.") from exc
        if not rows:
            return []

        headers = rows[0]
        pages: list[Page] = [{
            "page_number": 1,
            "text": f"CSV File Summary\nColumns ({len(headers)}): {', '.join(headers)}\n"
                    f"Total rows: {len(rows) - 1}",
            "metadata": {"parser": "csv", "type": "summary", "columns": headers},
        }]
        step = 50
        for i, start in enumerate(range(1, len(rows), step), start=2):
            block = rows[start:start + step]
            lines = ["\t".join(headers)] + ["\t".join(str(c) for c in row) for row in block]
            pages.append({
                "page_number": i, "text": "\n".join(lines),
                "metadata": {"parser": "csv", "type": "data",
                             "row_start": start, "row_end": min(start + step - 1, len(rows) - 1)},
            })
        return pages

    def _parse_excel(self, path: Path) -> list[Page]:
        try:
            import pandas as pd

            xl = pd.ExcelFile(str(path))
            pages: list[Page] = []
            num = 1
            step = 50
            for sheet in xl.sheet_names:
                df = xl.parse(sheet)
                pages.append({
                    "page_number": num,
                    "text": f"Sheet: {sheet}\nColumns ({len(df.columns)}): "
                            f"{', '.join(str(c) for c in df.columns)}\nRows: {len(df)}",
                    "metadata": {"parser": "pandas", "sheet": sheet, "type": "summary"},
                })
                num += 1
                for start in range(0, len(df), step):
                    block = df.iloc[start:start + step]
                    pages.append({
                        "page_number": num, "text": block.to_csv(index=False),
                        "metadata": {"parser": "pandas", "sheet": sheet, "type": "data",
                                     "row_start": start, "row_end": min(start + step - 1, len(df) - 1)},
                    })
                    num += 1
            return pages
        except ImportError as exc:
            raise DocumentParseError("Excel support is not installed.") from exc
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The Excel file could not be read.") from exc

    def _parse_text(self, path: Path) -> list[Page]:
        return self._paginate_text(self._read_text(path), "text")

    @staticmethod
    def _read_text(path: Path) -> str:
        try:
            raw = path.read_bytes()
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The file could not be read.") from exc
        return _decode(raw)

    @staticmethod
    def _paginate_text(text: str, parser: str) -> list[Page]:
        sections = [s.strip() for s in re.split(r"\n{2,}", text) if s.strip()]

        pages: list[Page] = []
        current: list[str] = []
        num = 1
        for section in sections:
            current.append(section)
            if sum(len(s) for s in current) >= 800:
                pages.append({"page_number": num, "text": "\n\n".join(current),
                              "metadata": {"parser": parser}})
                num += 1
                current = []
        if current:
            pages.append({"page_number": num, "text": "\n\n".join(current),
                          "metadata": {"parser": parser}})
        return pages

    def _parse_json(self, path: Path) -> list[Page]:
        text = self._read_text(path)
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError as exc:
            raise DocumentParseError("The JSON file is not valid JSON.") from exc
        # one line per element → group lines (not blank-line sections) into ~800-char pages
        pages: list[Page] = []
        current: list[str] = []
        size = 0
        for line in text.splitlines():
            current.append(line)
            size += len(line) + 1
            if size >= 800:
                pages.append({"page_number": len(pages) + 1, "text": "\n".join(current),
                              "metadata": {"parser": "json"}})
                current, size = [], 0
        if current:
            pages.append({"page_number": len(pages) + 1, "text": "\n".join(current),
                          "metadata": {"parser": "json"}})
        return pages

    def _parse_html(self, path: Path) -> list[Page]:
        text, parser_name = self._extract_article(self._read_text(path))
        if not text or not text.strip():
            raise DocumentParseError("No readable text found in the HTML file.")
        return self._paginate_text(text, parser_name)

    def _parse_rtf(self, path: Path) -> list[Page]:
        try:
            from striprtf.striprtf import rtf_to_text

            text = rtf_to_text(self._read_text(path), errors="ignore")
        except ImportError as exc:
            raise DocumentParseError("RTF support is not installed.") from exc
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The RTF file could not be read.") from exc
        return self._paginate_text(text, "rtf")

    def _parse_odt(self, path: Path) -> list[Page]:
        try:
            from lxml import etree

            with zipfile.ZipFile(str(path)) as zf:
                info = zf.getinfo("content.xml")
                if info.file_size > 50 * 1024 * 1024:
                    raise DocumentParseError("The ODT file is too large to read.")
                root = etree.fromstring(zf.read("content.xml"),
                                        etree.XMLParser(resolve_entities=False, no_network=True))
            ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
            lines = ["".join(el.itertext()).strip()
                     for el in root.iter(f"{ns}h", f"{ns}p") if "".join(el.itertext()).strip()]
        except DocumentParseError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The ODT file could not be read.") from exc
        return self._paginate_text("\n\n".join(lines), "odt")

    def _parse_pptx(self, path: Path) -> list[Page]:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise DocumentParseError("PPTX support is not installed.") from exc
        try:
            prs = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError("The PPTX file could not be read.") from exc

        pages: list[Page] = []
        for num, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                self._pptx_shape_text(shape, parts)
            try:
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        parts.append(f"Notes: {notes}")
            except Exception:  # noqa: BLE001 — notes are optional
                pass
            text = "\n".join(p for p in parts if p)
            if text.strip():
                pages.append({"page_number": num, "text": text,
                              "metadata": {"parser": "python-pptx", "type": "slide"}})
        return pages

    def _pptx_shape_text(self, shape: Any, parts: list[str]) -> None:
        if getattr(shape, "shapes", None) is not None:          # group shape
            for child in shape.shapes:
                self._pptx_shape_text(child, parts)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                parts.append(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                parts.append("\t".join(cell.text.strip() for cell in row.cells))

    @staticmethod
    def _split_into_chunks(text: str, chunk_size: int = 1200) -> list[str]:
        """Split text into roughly equal chunks at paragraph boundaries."""
        chunks: list[str] = []
        current: list[str] = []
        length = 0
        for para in re.split(r"\n{2,}", text):
            if length + len(para) > chunk_size and current:
                chunks.append("\n\n".join(current))
                current, length = [para], len(para)
            else:
                current.append(para)
                length += len(para)
        if current:
            chunks.append("\n\n".join(current))
        return chunks or [text[:chunk_size]]


# ── helpers ──────────────────────────────────────────────────────────────────

def _decode(raw: bytes) -> str:
    """Decode text bytes: UTF-8 (with BOM) first, then detect, never raise."""
    for enc in ("utf-8-sig",):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="replace")
    try:
        return raw.decode("cp1252")                  # the common legacy Western encoding
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(raw).best()
        if best is not None:
            return str(best)
    except Exception as exc:  # noqa: BLE001 — fall through to a lossy decode
        logger.debug("charset detection failed: %s", exc)
    return raw.decode("cp1252", errors="replace")


def _get_ocr() -> Any:
    """Lazily create the shared ONNX OCR engine; None when unavailable."""
    global _OCR_ENGINE, _OCR_FAILED
    if _OCR_ENGINE is not None or _OCR_FAILED:
        return _OCR_ENGINE
    with _OCR_LOCK:
        if _OCR_ENGINE is None and not _OCR_FAILED:
            try:
                from rapidocr_onnxruntime import RapidOCR

                _OCR_ENGINE = RapidOCR(rec_batch_num=1)   # batching pads wide lines: slower
            except Exception as exc:  # noqa: BLE001
                _OCR_FAILED = True
                logger.warning("OCR engine unavailable: %s", exc)
    return _OCR_ENGINE


def ocr_available() -> bool:
    return _get_ocr() is not None


def _limit_side(image: Any) -> Any:
    longest = max(image.size)
    if longest <= MAX_IMAGE_SIDE:
        return image
    ratio = MAX_IMAGE_SIDE / longest
    return image.resize((max(1, int(image.width * ratio)), max(1, int(image.height * ratio))))


def _ocr_array(engine: Any, array: Any) -> str:
    """Run OCR on an RGB array and return text in reading order (one box per line)."""
    with _OCR_LOCK:                       # the ONNX session is not guaranteed re-entrant
        result, _ = engine(array, use_cls=False)   # upright documents: skip the angle model
    if not result:
        return ""
    boxes = sorted(result, key=lambda r: (round(r[0][0][1] / 12), r[0][0][0]))
    return "\n".join(str(r[1]).strip() for r in boxes if str(r[1]).strip())
