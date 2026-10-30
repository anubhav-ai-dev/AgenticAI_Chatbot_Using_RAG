"""
Structured document attached to an assistant turn.

The model fills JSON. This module validates it and, only when a download is
requested, renders Markdown, plain text, Excel, Word, or PDF. File bytes are
never stored.
"""

from __future__ import annotations

import json
import logging
import re
from io import BytesIO
from typing import Any
from urllib.parse import quote

logger = logging.getLogger(__name__)

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# Preview opens only when the user names a downloadable format (or restyles).
_FILE_KEYWORD = re.compile(
    r"("
    r"\.(?:pdf|docx|xlsx|xlxs|xls|md|txt|csv)\b"
    r"|\btext\s+files?\b"
    r"|\b(?:pdf|excel|xlsx|xlxs|xls|csv|docx|markdown|document|word)\b"
    r")",
    re.IGNORECASE,
)
_STYLE_WORDS = re.compile(
    r"\b(?:apa|mla|chicago|landscape|portrait|cover\s*page|one[-\s]?pager|"
    r"blue|red|green|brand|colour|color|font|headings?|margin|spacing|"
    r"letterhead|formal|casual|executive|larger|smaller)\b",
    re.IGNORECASE,
)
_RESTYLE_VERB = re.compile(
    r"\b(?:make|change|use|apply|switch|set|increase|decrease|restyle|reformat)\b",
    re.IGNORECASE,
)
_HEADING = re.compile(r"(?m)^#{1,6}\s+\S")
_UNSAFE_NAME = re.compile(r"[^\w.\- ]+", re.UNICODE)
_FILENAME_MAX_WORDS = 4
_FILENAME_STOP = {
    "a", "an", "and", "are", "as", "at", "be", "been", "being", "by", "for",
    "from", "how", "in", "into", "is", "it", "its", "of", "on", "or", "our",
    "please", "that", "the", "their", "these", "this", "those", "to", "was",
    "were", "what", "when", "where", "which", "who", "why", "with", "your",
    "csv", "document", "docx", "excel", "file", "markdown", "pdf", "txt",
    "word", "xlsx",
}
_FENCE = re.compile(r"^```(?:json)?\s*|\s*```$", re.IGNORECASE)
_INLINE_LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_INLINE_CODE = re.compile(r"`([^`]+)`")
_INLINE_BOLD = re.compile(r"(\*\*|__)(.+?)\1")
_INLINE_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)|(?<!_)_(?!_)(.+?)(?<!_)_(?!_)")
_INLINE_RUN = re.compile(
    r"(\*\*(.+?)\*\*|__(.+?)__|(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)|"
    r"(?<!_)_(?!_)(.+?)(?<!_)_(?!_)|`([^`]+)`|\[([^\]]+)\]\(([^)]+)\))"
)

_SYSTEM = (
    "Turn the assistant answer into one JSON document the user can download. "
    "Return only a JSON object, with no markdown fences, using this shape: "
    '{"title":"short title","blocks":['
    '{"type":"heading","text":"..."},'
    '{"type":"paragraph","text":"..."},'
    '{"type":"table","columns":["A"],"rows":[["1"]]}'
    "]}. "
    "title is 1-120 characters. Use 1-200 blocks. "
    "heading and paragraph text are plain strings of at most 8000 characters. "
    "A table has 1-20 column names and at most 5000 rows. "
    "Every cell is a string of at most 1000 characters. "
    "Use heading, paragraph, list, code, and table only. Do not include HTML. "
    "Cover the answer. Do not invent sections that are not in the answer."
)

FORMATS = ("md", "txt", "xlsx", "docx", "pdf")
_MEDIA = {
    "md": "text/markdown; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


class ArtifactError(Exception):
    """The document was empty, over the caps, or not the expected JSON."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


FILE_ANSWER_NOTE = (
    "The user asked for a downloadable file. You CAN produce it. "
    "Never say you cannot generate, send, attach, or export a PDF, Word, Excel, Markdown, or text file. "
    "Do not mention copying into Word, Google Docs, Notepad, or a printer. "
    "Write the document body in markdown now. Title it from the topic, not the format word. "
    "Use a professional layout: a clear title, a short opening, hierarchical headings in title case, "
    "scannable lists, tables for comparisons, **bold** where it helps. "
    "The application converts your markdown into the download and opens the preview."
)
FILE_SAVE_NOTE = (
    "The user asked to save the previous answer as a file. You CAN do that. "
    "Never say you cannot generate or send a file. Reply in one or two sentences naming what "
    "the file contains. Do not repeat the full document. Do not explain how to export. "
    "The application attaches the previous answer as the download and opens the preview."
)
FILE_RESTYLE_NOTE = (
    "The user wants the previous document restyled. Rewrite it with their formatting changes "
    "and keep the same facts. Do not explain how to export a file. The application attaches "
    "the download."
)

_EXPORT_HELP = re.compile(
    r"("
    r"print to pdf|microsoft print|export as pdf|notepad|textedit|google docs|"
    r"online converter|copy into any text editor|copy-paste|"
    r"file\s*[→>-]+\s*(print|download)|"
    r"cannot (directly )?(generate|create|send|export|attach|produce)|"
    r"i cannot (directly )?(generate|create|send)|"
    r"copy it into a document|save it as a pdf"
    r")",
    re.IGNORECASE,
)
_FORMAT_WORD = re.compile(
    r"\.(?:pdf|docx|xlsx|xlxs|xls|md|txt|csv)\b|"
    r"\btext\s+files?\b|"
    r"\b(?:pdf|docx|xlsx|xlxs|xls|csv|excel|spreadsheet|markdown|document|word|"
    r"plain\s+text|word\s+doc(?:ument)?)\b",
    re.IGNORECASE,
)
_NAVY_RGB = (0x12, 0x1A, 0x45)
_ACCENT_RGB = (0x4F, 0x46, 0xE5)
_MUTED_RGB = (0x4A, 0x56, 0x7F)
_CODE_BG_RGB = (0xF3, 0xF5, 0xFB)
_WHITE_RGB = (255, 255, 255)
_NAVY_HEX = "121A45"
_ACCENT_HEX = "4F46E5"
_BORDER_HEX = "E4E7F3"
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)*\|?\s*$")


def wants_file(user_text: str) -> bool:
    """True when the user named a downloadable format, or asked to restyle the last file."""
    text = (user_text or "").strip()
    if not text:
        return False
    if is_restyle(text):
        return True
    return bool(_FILE_KEYWORD.search(text))


def is_restyle(user_text: str) -> bool:
    """True when a short follow-up asks to change the previous document's formatting."""
    text = (user_text or "").strip()
    if not text or len(text.split()) > 16:
        return False
    return bool(_STYLE_WORDS.search(text) and _RESTYLE_VERB.search(text))


def style_note(user_text: str) -> str:
    if _STYLE_WORDS.search(user_text or ""):
        return ("Apply the user's formatting instructions from their message. "
                "Keep the result professional.")
    return ""


_BARE_LEFTOVER = re.compile(
    r"^(?:generate|create|make|export|download|prepare|save|convert|write|draft|preview)?$",
    re.IGNORECASE,
)
_LEAD = re.compile(r"^([A-Z][^.\n]{2,60})\.\s+(\S.+)$", re.DOTALL)
_LABEL = re.compile(r"^([A-Z][^:\n]{2,40}):\s*(\S.+)$", re.DOTALL)
_GENERIC_TITLE = {
    "document", "generate", "create", "make", "download", "export", "write", "draft",
    "pdf", "excel", "word", "markdown", "csv", "txt", "text", "file", "preview",
}


def is_bare_save(user_text: str) -> bool:
    """True when the user only asked to save the last answer, as in 'i want a pdf'."""
    text = (user_text or "").strip()
    if not wants_file(text) or len(text.split()) > 8:
        return False
    leftover = _FILE_KEYWORD.sub(" ", text)
    leftover = re.sub(
        r"\b(?:please|can|you|i|we|they|a|an|the|this|that|it|me|my|as|to|into|for|of|"
        r"in|on|and|file|now|just|notes|report|proposal|spreadsheet|plain|preview|"
        r"download|export|save|generate|create|make|prepare|convert|write|draft|"
        r"want|need|like|get|give|send|would|could|should|some|any|one|here|there|"
        r"also|really)\b",
        " ", leftover, flags=re.IGNORECASE)
    leftover = re.sub(r"[.\-]+", " ", leftover)
    leftover = re.sub(r"\s+", " ", leftover).strip()
    return leftover == "" or bool(_BARE_LEFTOVER.match(leftover))


def topic_title(user_text: str, history: list[dict[str, str]] | None = None) -> str:
    """Title for a file. A bare 'generate pdf' uses the question that came before it."""
    if is_bare_save(user_text) or is_restyle(user_text):
        for message in reversed(history or []):
            if message.get("role") != "user":
                continue
            content = (message.get("content") or "").strip()
            if content and not is_bare_save(content):
                source = fallback_title(content) if wants_file(content) else content
                return _nice_title(source)
        return ""
    return _nice_title(fallback_title(user_text))


def should_author_document(user_text: str, answer: str) -> bool:
    """True when this turn should produce a downloadable document.

    Opens when the user names a format (pdf, excel, document, word, markdown,
    csv, text file, or a matching extension) or asks to restyle the last file.
    """
    text = (answer or "").strip()
    if not text:
        return False
    if len(text) <= 400 and text.endswith("?"):
        return False
    return wants_file(user_text)


def document_source(user_text: str, answer: str, history: list[dict[str, str]] | None = None) -> str:
    """Use the previous answer when this turn only asks to save what was just written."""
    previous = ""
    for message in reversed(history or []):
        if message.get("role") == "assistant" and (message.get("content") or "").strip():
            previous = message["content"].strip()
            break
    current = (answer or "").strip()
    if previous and _use_previous(user_text, current, previous):
        return previous
    return current


def document_from_markdown(markdown: str, *, fallback_title: str = "") -> dict[str, Any]:
    """Turn markdown the model already wrote into the validated document. No extra model call."""
    blocks = _fit_blocks(_promote_sections(_parse_markdown(markdown or "")))
    if not blocks:
        plain = _clean(markdown or "").strip() or _clip_title(fallback_title or "Document")
        blocks = _fit_blocks([{"type": "paragraph", "text": plain}])
    title, blocks = _choose_title(blocks, fallback_title)
    if not blocks:
        blocks = [{"type": "paragraph", "text": title}]
    return validate_document({"title": title, "blocks": blocks})


def preferred_format(user_text: str) -> str:
    """Which download the user named. Markdown when they did not name one."""
    text = user_text or ""
    if re.search(r"\bpdf\b|\.pdf\b", text, re.IGNORECASE):
        return "pdf"
    if re.search(r"\b(?:docx|word)\b|\.docx\b", text, re.IGNORECASE):
        return "docx"
    if re.search(r"\b(?:xlsx|xlxs|xls|excel|spreadsheet|csv)\b|\.(?:xlsx|xlxs|xls|csv)\b",
                 text, re.IGNORECASE):
        return "xlsx"
    if re.search(r"\b(?:plain\s+text|text\s+files?)\b|\.txt\b", text, re.IGNORECASE):
        return "txt"
    if re.search(r"\bmarkdown\b|\.md\b", text, re.IGNORECASE):
        return "md"
    return "md"


def fallback_title(user_text: str) -> str:
    text = _FORMAT_WORD.sub(" ", user_text or "")
    text = re.sub(r"\s+", " ", text).strip(" .")
    return _clip_title(text or "Document")


def artifact_messages(user_text: str, answer: str) -> list[dict[str, str]]:
    return [
        {"role": "system", "content": _SYSTEM},
        {"role": "user", "content": (
            f"User request:\n{(user_text or '')[:4000]}\n\n"
            f"Answer:\n{(answer or '')[:24000]}\n\nJSON document:"
        )},
    ]


def parse_model_document(text: str) -> Any:
    raw = _FENCE.sub("", (text or "").strip()).strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start >= 0 and end > start:
            try:
                return json.loads(raw[start:end + 1])
            except json.JSONDecodeError:
                pass
        raise ArtifactError("The document was not valid.")


def validate_document(raw: Any) -> dict[str, Any]:
    """Drop unknown fields, strip control characters, and reject empty or over-cap documents."""
    if not isinstance(raw, dict):
        raise ArtifactError("The document was not valid.")
    title = _plain(raw.get("title"), "The document was not valid.")
    if not title:
        raise ArtifactError("The document is empty.")
    if len(title) > 120:
        raise ArtifactError("The document title is too long.")
    blocks = raw.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        raise ArtifactError("The document is empty.")
    if len(blocks) > 200:
        raise ArtifactError("The document is too large to save.")
    clean: list[dict[str, Any]] = []
    for block in blocks:
        clean.append(_block(block))
    return {"title": title, "blocks": clean}


def render_download(doc: dict[str, Any], fmt: str) -> tuple[bytes, str, str]:
    """Build file bytes for one format. `doc` must already be validated."""
    if fmt not in _MEDIA:
        raise ArtifactError("That download format is not supported.")
    try:
        if fmt == "md":
            data = render_markdown(doc).encode("utf-8")
        elif fmt == "txt":
            data = render_text(doc).encode("utf-8")
        elif fmt == "xlsx":
            data = render_xlsx(doc)
        elif fmt == "docx":
            data = render_docx(doc)
        else:
            data = render_pdf(doc)
    except ArtifactError:
        raise
    except Exception as exc:
        logger.info("artifact render failed (%s)", type(exc).__name__)
        raise ArtifactError("The file could not be created.") from exc
    return data, _MEDIA[fmt], safe_filename(doc["title"], fmt)


def content_disposition(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1] if "." in filename else "bin"
    ascii_name = filename.encode("ascii", "ignore").decode("ascii").replace('"', "").replace("\\", "")
    if not ascii_name or ascii_name.startswith(".") or "." not in ascii_name:
        ascii_name = f"document.{ext}"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


def safe_filename(title: str, ext: str) -> str:
    """Short download name from the document heading: at most 4 words, joined with _."""
    cleaned = _UNSAFE_NAME.sub(" ", title or "")
    cleaned = re.sub(r"[\s_\-]+", " ", cleaned).strip(" .")
    words: list[str] = []
    for token in cleaned.split():
        if token.lower() in _FILENAME_STOP:
            continue
        if token.isupper() or (token.isalnum() and token[:1].isdigit()):
            words.append(token)
        else:
            words.append(token[:1].upper() + token[1:])
        if len(words) >= _FILENAME_MAX_WORDS:
            break
    name = "_".join(words)[:80].strip("._") or "Document"
    return f"{name}.{ext}"


def strip_markdown_inline(text: str) -> str:
    """Remove markdown emphasis so PDF/Word/Excel/plain text never show leftover asterisks."""
    value = text or ""
    value = _INLINE_LINK.sub(r"\1", value)
    value = _INLINE_CODE.sub(r"\1", value)
    value = _INLINE_BOLD.sub(r"\2", value)
    value = _INLINE_ITALIC.sub(lambda m: m.group(1) or m.group(2) or "", value)
    return value


def render_markdown(doc: dict[str, Any]) -> str:
    lines = [f"# {_one_line(doc['title'])}", ""]
    for block in doc["blocks"]:
        kind = block["type"]
        if kind == "heading":
            marks = "#" * _heading_marks(block)
            lines.extend((f"{marks} {_one_line(block['text'])}", ""))
        elif kind == "paragraph":
            lines.extend((block["text"], ""))
        elif kind == "list":
            for i, item in enumerate(block.get("items") or [], 1):
                prefix = f"{i}. " if block.get("ordered") else "- "
                lines.append(prefix + item)
            lines.append("")
        elif kind == "code":
            lang = (block.get("language") or "").strip()
            lines.append(f"```{lang}".rstrip())
            lines.append(block.get("text") or "")
            lines.extend(("```", ""))
        else:
            cols = [_md_cell(c) for c in block["columns"]]
            lines.append("| " + " | ".join(cols) + " |")
            lines.append("| " + " | ".join("---" for _ in cols) + " |")
            width = len(cols)
            for row in block["rows"]:
                cells = [_md_cell(c) for c in _pad(row, width)]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_text(doc: dict[str, Any]) -> str:
    title = strip_markdown_inline(doc["title"])
    lines = [title, "=" * min(48, max(3, len(title))), ""]
    for block in doc["blocks"]:
        kind = block["type"]
        if kind == "heading":
            heading = strip_markdown_inline(block["text"])
            lines.extend((heading, "-" * min(40, max(3, len(heading))), ""))
        elif kind == "paragraph":
            lines.extend((strip_markdown_inline(block["text"]), ""))
        elif kind == "list":
            for i, item in enumerate(block.get("items") or [], 1):
                prefix = f"  {i}. " if block.get("ordered") else "  • "
                lines.append(prefix + strip_markdown_inline(item))
            lines.append("")
        elif kind == "code":
            for row in (block.get("text") or "").splitlines() or [""]:
                lines.append("    " + row)
            lines.append("")
        else:
            lines.append("  " + "\t".join(strip_markdown_inline(c) for c in block["columns"]))
            width = len(block["columns"])
            for row in block["rows"]:
                lines.append("  " + "\t".join(strip_markdown_inline(c) for c in _pad(row, width)))
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def render_xlsx(doc: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = _sheet_name(doc["title"])
    title_font = Font(name="Calibri", bold=True, size=18, color=_ACCENT_HEX, underline="single")
    heading_font = Font(name="Calibri", bold=True, size=13, color=_NAVY_HEX, underline="single")
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, color=_NAVY_HEX)
    header_fill = PatternFill("solid", fgColor=_ACCENT_HEX)
    title_fill = PatternFill("solid", fgColor="EEF0FF")
    thin = Border(
        left=Side(style="thin", color=_BORDER_HEX),
        right=Side(style="thin", color=_BORDER_HEX),
        top=Side(style="thin", color=_BORDER_HEX),
        bottom=Side(style="thin", color=_BORDER_HEX),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    indent = Alignment(wrap_text=True, vertical="top", indent=1)
    row_idx = 1
    tables: list[tuple[str, dict[str, Any]]] = []
    last_heading = "Table"
    widths: dict[int, int] = {}

    def _put(r: int, c: int, value: str, *, kind: str = "body") -> None:
        cell = sheet.cell(r, c, value)
        cell.alignment = indent if kind == "list" else wrap
        if kind == "title":
            cell.font = title_font
            cell.fill = title_fill
        elif kind == "heading":
            cell.font = heading_font
        elif kind == "header":
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = thin
        else:
            cell.font = body_font
            if kind == "table":
                cell.border = thin
        first = value.split("\n", 1)[0]
        widths[c] = max(widths.get(c, 12), min(48, max(12, len(first) + 2)))

    _put(row_idx, 1, strip_markdown_inline(doc["title"]), kind="title")
    sheet.row_dimensions[row_idx].height = 24
    row_idx += 2

    for block in doc["blocks"]:
        kind = block["type"]
        if kind == "heading":
            last_heading = strip_markdown_inline(block["text"]) or last_heading
            _put(row_idx, 1, last_heading, kind="heading")
            row_idx += 2
            continue
        if kind == "paragraph":
            _put(row_idx, 1, strip_markdown_inline(block["text"]))
            row_idx += 2
            continue
        if kind == "list":
            for i, item in enumerate(block.get("items") or [], 1):
                prefix = f"{i}. " if block.get("ordered") else "• "
                _put(row_idx, 1, prefix + strip_markdown_inline(item), kind="list")
                row_idx += 1
            row_idx += 1
            continue
        if kind == "code":
            _put(row_idx, 1, block.get("text") or "")
            row_idx += 2
            continue
        tables.append((last_heading, block))
        for col, name in enumerate(block["columns"], start=1):
            _put(row_idx, col, strip_markdown_inline(name), kind="header")
        row_idx += 1
        width = len(block["columns"])
        for values in block["rows"]:
            for col, value in enumerate(_pad(values, width), start=1):
                _put(row_idx, col, strip_markdown_inline(value), kind="table")
            row_idx += 1
        row_idx += 1
    for col, size in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = size
    used = {sheet.title.lower()}
    if len(tables) >= 2:
        for heading, table in tables:
            extra = book.create_sheet(_unique_sheet(heading, used))
            used.add(extra.title.lower())
            for col, name in enumerate(table["columns"], start=1):
                cell = extra.cell(1, col, strip_markdown_inline(name))
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                cell.border = thin
            width = len(table["columns"])
            for r, values in enumerate(table["rows"], start=2):
                for col, value in enumerate(_pad(values, width), start=1):
                    cell = extra.cell(r, col, strip_markdown_inline(value))
                    cell.font = body_font
                    cell.alignment = wrap
                    cell.border = thin
            for col in range(1, width + 1):
                extra.column_dimensions[get_column_letter(col)].width = 18
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def render_docx(doc: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    out = Document()
    try:
        normal = out.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(*_NAVY_RGB)
    except KeyError:
        pass
    for section in out.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)
    title = out.add_heading(strip_markdown_inline(doc["title"]), level=1)
    _style_docx_heading(title, 22, _ACCENT_RGB, underline=True)
    for block in doc["blocks"]:
        kind = block["type"]
        if kind == "heading":
            level = 2 if _heading_level(block) <= 2 else 3
            heading = out.add_heading(strip_markdown_inline(block["text"]), level=level)
            _style_docx_heading(heading, 16 if level == 2 else 13, _NAVY_RGB, underline=level == 2)
        elif kind == "paragraph":
            _style_docx_body(out.add_paragraph(), block["text"])
        elif kind == "list":
            style = "List Number" if block.get("ordered") else "List Bullet"
            for item in block.get("items") or []:
                try:
                    para = out.add_paragraph(style=style)
                except KeyError:
                    para = out.add_paragraph()
                _style_docx_body(para, item)
        elif kind == "code":
            para = out.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(10)
            run = para.add_run(block.get("text") or "")
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(*_NAVY_RGB)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:eastAsia"), "Courier New")
        else:
            width = len(block["columns"])
            table = out.add_table(rows=1, cols=width)
            try:
                table.style = "Table Grid"
            except KeyError:
                pass
            for idx, name in enumerate(block["columns"]):
                _docx_cell(table.rows[0].cells[idx], strip_markdown_inline(name), header=True)
            for values in block["rows"]:
                cells = table.add_row().cells
                for idx, value in enumerate(_pad(values, width)):
                    _docx_cell(cells[idx], strip_markdown_inline(value))
    buf = BytesIO()
    out.save(buf)
    return buf.getvalue()


def render_pdf(doc: dict[str, Any]) -> bytes:
    from fpdf import FPDF

    class _DocPDF(FPDF):
        def footer(self) -> None:
            self.set_y(-12)
            self.set_font("Helvetica", "", 8)
            self.set_text_color(*_MUTED_RGB)
            self.cell(0, 8, str(self.page_no()), align="C")
            self.set_text_color(*_NAVY_RGB)

    pdf = _DocPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_left_margin(18)
    pdf.set_right_margin(18)
    pdf.add_page()

    def _line(text: str, height: float, align: str = "L", fill: bool = False, md: bool = True) -> None:
        payload = _pdf_text(text)
        try:
            pdf.multi_cell(0, height, payload, align=align, fill=fill, markdown=md,
                           new_x="LMARGIN", new_y="NEXT")
        except Exception:
            pdf.multi_cell(0, height, _pdf_text(strip_markdown_inline(text)), align=align,
                           fill=fill, new_x="LMARGIN", new_y="NEXT")

    def _heading(text: str, size: float, color: tuple[int, int, int], underline: bool) -> None:
        pdf.set_font("Helvetica", "B", size)
        pdf.set_text_color(*color)
        _line(strip_markdown_inline(text), size * 0.5, md=False)
        if underline:
            y = pdf.get_y()
            pdf.set_draw_color(*_ACCENT_RGB)
            pdf.set_line_width(0.55 if size >= 16 else 0.35)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(2)
        pdf.set_text_color(*_NAVY_RGB)

    pdf.set_fill_color(*_ACCENT_RGB)
    pdf.set_text_color(*_WHITE_RGB)
    pdf.set_font("Helvetica", "B", 18)
    _line(strip_markdown_inline(doc["title"]), 11, fill=True, md=False)
    pdf.set_text_color(*_NAVY_RGB)
    pdf.ln(5)
    for block in doc["blocks"]:
        kind = block["type"]
        if kind == "heading":
            pdf.ln(2)
            level = _heading_level(block)
            _heading(block["text"], 14 if level <= 2 else 12, _ACCENT_RGB if level <= 2 else _NAVY_RGB,
                     underline=level <= 2)
        elif kind == "paragraph":
            note = _one_line(strip_markdown_inline(block["text"])).lower().startswith(("caveat:", "evidence"))
            pdf.set_font("Helvetica", "", 11)
            pdf.set_text_color(*_NAVY_RGB)
            if note:
                pdf.set_fill_color(255, 244, 204)
            _line(block["text"], 6, fill=note)
            pdf.ln(1)
        elif kind == "list":
            pdf.set_font("Helvetica", "", 11)
            pdf.set_text_color(*_NAVY_RGB)
            left = pdf.l_margin
            pdf.set_left_margin(left + 8)
            pdf.set_x(left + 8)
            for i, item in enumerate(block.get("items") or [], 1):
                prefix = f"{i}. " if block.get("ordered") else "- "
                _line(prefix + item, 6)
            pdf.set_left_margin(left)
            pdf.set_x(left)
            pdf.ln(1)
        elif kind == "code":
            pdf.set_font("Courier", "", 9)
            pdf.set_fill_color(*_CODE_BG_RGB)
            pdf.set_text_color(*_NAVY_RGB)
            _line(block.get("text") or "", 5, fill=True, md=False)
            pdf.ln(1)
        else:
            _pdf_table(pdf, block["columns"], block["rows"])
            pdf.ln(1)
    payload = pdf.output()
    return bytes(payload)


def _block(raw: Any) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ArtifactError("The document was not valid.")
    kind = raw.get("type")
    if kind == "heading":
        text = _plain(raw.get("text"), "The document was not valid.")
        if not text:
            raise ArtifactError("A section is empty.")
        if len(text) > 8000:
            raise ArtifactError("A section is too long to save.")
        return {"type": "heading", "text": text, "level": _heading_level(raw)}
    if kind == "paragraph":
        text = _plain(raw.get("text"), "The document was not valid.")
        if not text:
            raise ArtifactError("A section is empty.")
        if len(text) > 8000:
            raise ArtifactError("A section is too long to save.")
        return {"type": "paragraph", "text": text}
    if kind == "list":
        return _list_block(raw)
    if kind == "code":
        text = _plain(raw.get("text"), "The document was not valid.")
        if not text:
            raise ArtifactError("A section is empty.")
        if len(text) > 8000:
            raise ArtifactError("A section is too long to save.")
        lang = _clean(str(raw.get("language") or "")).strip()[:32]
        return {"type": "code", "language": lang, "text": text}
    if kind == "table":
        return _table(raw)
    raise ArtifactError("The document was not valid.")


def _list_block(raw: dict[str, Any]) -> dict[str, Any]:
    items_raw = raw.get("items")
    if not isinstance(items_raw, list) or not items_raw:
        raise ArtifactError("A section is empty.")
    if len(items_raw) > 200:
        raise ArtifactError("The document is too large to save.")
    items = []
    for item in items_raw:
        text = _plain(item, "The document was not valid.")
        if not text:
            continue
        if len(text) > 8000:
            raise ArtifactError("A section is too long to save.")
        items.append(text)
    if not items:
        raise ArtifactError("A section is empty.")
    return {"type": "list", "ordered": bool(raw.get("ordered")), "items": items}


def _table(raw: dict[str, Any]) -> dict[str, Any]:
    columns_raw = raw.get("columns")
    rows_raw = raw.get("rows")
    if not isinstance(columns_raw, list) or not isinstance(rows_raw, list):
        raise ArtifactError("The document was not valid.")
    if not columns_raw:
        raise ArtifactError("The table has no columns.")
    if len(columns_raw) > 20:
        raise ArtifactError("The table has too many columns.")
    if len(rows_raw) > 5000:
        raise ArtifactError("The table has too many rows.")
    columns = [_cell(name) for name in columns_raw]
    width = len(columns)
    rows: list[list[str]] = []
    for row in rows_raw:
        if not isinstance(row, list):
            raise ArtifactError("The document was not valid.")
        if len(row) > width:
            raise ArtifactError("A table row has too many cells.")
        rows.append([_cell(value) for value in row] + [""] * (width - len(row)))
    return {"type": "table", "columns": columns, "rows": rows}


def _plain(value: Any, invalid: str) -> str:
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ArtifactError(invalid)
    return _clean(str(value)).strip()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ArtifactError("The document was not valid.")
    text = _clean(str(value))
    if len(text) > 1000:
        raise ArtifactError("A table cell is too long.")
    return text


def _clean(value: str) -> str:
    return _CONTROL.sub("", value).replace("\r\n", "\n").replace("\r", "\n")


def _one_line(value: str) -> str:
    return value.replace("\n", " ").strip()


def _md_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def _pad(row: list[str], width: int) -> list[str]:
    cells = list(row[:width])
    if len(cells) < width:
        cells.extend([""] * (width - len(cells)))
    return cells


def _pdf_text(value: str) -> str:
    cleaned = value.encode("latin-1", "replace").decode("latin-1")
    parts: list[str] = []
    for token in cleaned.split(" "):
        if len(token) > 40:
            token = " ".join(token[i:i + 40] for i in range(0, len(token), 40))
        parts.append(token)
    return " ".join(parts) or " "


def _use_previous(user_text: str, answer: str, previous: str) -> bool:
    """A short 'generate pdf' should file the last answer, not a how-to-export reply."""
    if len(previous) < 80:
        return False
    if _EXPORT_HELP.search(answer or ""):
        return True
    if is_restyle(user_text):
        return True
    if not is_bare_save(user_text):
        return False
    # The new reply is already a full document, so keep it.
    if len(_HEADING.findall(answer or "")) >= 2 and len(answer or "") >= 400:
        return False
    return len(previous) > len(answer or "")


def _nice_title(value: str) -> str:
    text = re.sub(r"\s+", " ", value or "").strip(" .?:")
    text = re.sub(r"\bai\b", "AI", text, flags=re.IGNORECASE)
    if text[:1].islower():
        text = text[:1].upper() + text[1:]
    return _clip_title(text or "Document")


def _choose_title(blocks: list[dict[str, Any]], fallback_title: str) -> tuple[str, list[dict[str, Any]]]:
    """A markdown file keeps its opening heading. A prose save uses the user's question."""
    fallback = (fallback_title or "").strip()
    if blocks and blocks[0]["type"] == "heading":
        title = _clip_title(blocks[0]["text"])
        rest = blocks[1:]
        return title, rest or [{"type": "paragraph", "text": title}]
    if fallback and fallback.lower() not in _GENERIC_TITLE:
        return _clip_title(fallback), blocks
    heading = next((b["text"] for b in blocks if b["type"] == "heading"), "")
    return _clip_title(heading or fallback or "Document"), blocks


def _promote_sections(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Turn 'Multiagent systems. Gartner lists…' into a heading plus a paragraph."""
    if sum(1 for b in blocks if b["type"] == "heading") >= 2:
        return blocks
    out: list[dict[str, Any]] = []
    lead_at: list[int] = []
    for block in blocks:
        if block["type"] != "paragraph":
            out.append(block)
            continue
        text = (block.get("text") or "").strip()
        if not text or text.startswith("```") or _looks_like_code(text):
            out.append(block)
            continue
        bare = re.match(r"^([A-Z][^:\n]{2,40}):\s*$", text)
        if bare and len(bare.group(1).split()) <= 7:
            out.append({"type": "heading", "text": bare.group(1).strip(), "level": 2})
            continue
        lead = _LEAD.match(text)
        label = None if lead else _LABEL.match(text)
        match = lead or label
        title = match.group(1).strip() if match else ""
        if not match or len(title.split()) > 7:
            out.extend(_caveat_blocks(text))
            continue
        out.append({"type": "heading", "text": title, "level": 2})
        if lead:
            lead_at.append(len(out) - 1)
        rest = match.group(2).strip()
        if rest:
            out.extend(_caveat_blocks(rest))
    if len(lead_at) >= 3:
        for number, index in enumerate(lead_at, 1):
            out[index]["text"] = f"{number}. {out[index]['text']}"
    return out


def _caveat_blocks(text: str) -> list[dict[str, Any]]:
    parts = re.split(r"(?<=\s)(?=Caveat:)", text, maxsplit=1)
    return [{"type": "paragraph", "text": part.strip()} for part in parts if part.strip()]


def _looks_like_code(text: str) -> bool:
    first = text.split("\n", 1)[0]
    return first.endswith("{") or first.endswith(";") or bool(re.match(r"^(def |class |import |const |let |function )", first))


def _clip_title(value: str) -> str:
    text = _clean(value).replace("\n", " ").strip()
    return text[:120] or "Document"


def _parse_markdown(markdown: str) -> list[dict[str, Any]]:
    lines = _clean(markdown).split("\n")
    blocks: list[dict[str, Any]] = []
    paragraph: list[str] = []
    i = 0

    def flush() -> None:
        text = "\n".join(paragraph).strip()
        paragraph.clear()
        if text:
            blocks.append({"type": "paragraph", "text": text})

    while i < len(lines):
        line = lines[i]
        fence = re.match(r"^\s*```([\w#+.-]*)\s*$", line)
        if fence:
            flush()
            lang = fence.group(1).strip()
            body: list[str] = []
            i += 1
            while i < len(lines) and not re.match(r"^\s*```\s*$", lines[i]):
                body.append(lines[i])
                i += 1
            code = "\n".join(body).strip()
            if code:
                blocks.append({"type": "code", "language": lang, "text": code})
            i += 1
            continue
        ul = re.match(r"^\s{0,3}[-*+]\s+(.+)$", line)
        ol = re.match(r"^\s{0,3}(\d+)[.)]\s+(.+)$", line)
        if ul or ol:
            flush()
            ordered = bool(ol)
            items = [(ol.group(2) if ol else ul.group(1)).strip()]
            i += 1
            while i < len(lines):
                nxt = lines[i]
                more_ol = re.match(r"^\s{0,3}(\d+)[.)]\s+(.+)$", nxt)
                more_ul = re.match(r"^\s{0,3}[-*+]\s+(.+)$", nxt)
                if ordered and more_ol:
                    items.append(more_ol.group(2).strip())
                    i += 1
                    continue
                if not ordered and more_ul:
                    items.append(more_ul.group(1).strip())
                    i += 1
                    continue
                if nxt.startswith("  ") and nxt.strip() and not more_ol and not more_ul:
                    items[-1] = (items[-1] + " " + nxt.strip()).strip()
                    i += 1
                    continue
                break
            blocks.append({"type": "list", "ordered": ordered, "items": [x for x in items if x]})
            continue
        if i + 1 < len(lines) and "|" in line and _is_table_sep(lines[i + 1]):
            flush()
            columns = _split_row(line)
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i]))
                i += 1
            if columns:
                width = min(20, len(columns))
                blocks.append({
                    "type": "table",
                    "columns": columns[:width],
                    "rows": [row[:width] + [""] * max(0, width - len(row)) for row in rows[:5000]],
                })
            continue
        heading = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush()
            level = min(3, max(1, len(heading.group(1))))
            blocks.append({"type": "heading", "text": heading.group(2).strip(), "level": level})
            i += 1
            continue
        if not line.strip() or re.match(r"^\s*([-*_])\1{2,}\s*$", line):
            flush()
            i += 1
            continue
        paragraph.append(line.rstrip())
        i += 1
    flush()
    return blocks


def _is_table_sep(line: str) -> bool:
    # A plain "---" is a horizontal rule. A separator has at least one pipe.
    return "|" in line and bool(_TABLE_SEP.match(line))


def _split_row(line: str) -> list[str]:
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    return [cell.strip() for cell in raw.split("|")]


def _fit_blocks(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    fitted: list[dict[str, Any]] = []
    for block in blocks:
        kind = block.get("type")
        if kind == "table":
            columns = [_clip_cell(c) or " " for c in block["columns"][:20]]
            if not columns:
                continue
            rows = []
            for row in block["rows"][:5000]:
                cells = [_clip_cell(c) for c in row[:len(columns)]]
                cells.extend([""] * (len(columns) - len(cells)))
                rows.append(cells)
            fitted.append({"type": "table", "columns": columns, "rows": rows})
            continue
        if kind == "list":
            items = [_clip_cell(i) for i in (block.get("items") or []) if str(i).strip()]
            items = [i for i in items if i][:200]
            if items:
                fitted.append({"type": "list", "ordered": bool(block.get("ordered")), "items": items})
            continue
        if kind == "code":
            text = (block.get("text") or "").strip()[:8000]
            if text:
                fitted.append({
                    "type": "code",
                    "language": _clean(str(block.get("language") or "")).strip()[:32],
                    "text": text,
                })
            continue
        text = (block.get("text") or "").strip()
        out_kind = kind if kind in ("heading", "paragraph") else "paragraph"
        first = True
        while text:
            chunk = text[:8000]
            if len(text) > 8000:
                cut = max(chunk.rfind("\n"), chunk.rfind(" "))
                if cut >= 4000:
                    chunk = text[:cut]
            piece = chunk.strip()
            if piece:
                item: dict[str, Any] = {"type": out_kind if first else "paragraph", "text": piece}
                if item["type"] == "heading":
                    item["level"] = _heading_level(block)
                fitted.append(item)
                first = False
            text = text[len(chunk):].strip()
    if len(fitted) > 200:
        fitted = fitted[:199]
        fitted.append({
            "type": "paragraph",
            "text": "The rest of the answer was shortened so the file stays within its size limit.",
        })
    return fitted


def _clip_cell(value: Any) -> str:
    return _clean("" if value is None else str(value))[:1000]


def _sheet_name(title: str) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", " ", title).strip()
    return (cleaned[:31] or "Document")


def _unique_sheet(title: str, used: set[str]) -> str:
    base = _sheet_name(title)
    name = base
    n = 2
    while name.lower() in used:
        suffix = f" {n}"
        name = (base[: 31 - len(suffix)] + suffix).strip() or f"Sheet{n}"
        n += 1
    return name


def _heading_level(block: dict[str, Any]) -> int:
    level = block.get("level")
    if isinstance(level, bool) or not isinstance(level, int) or level < 1 or level > 3:
        return 2
    return level


def _heading_marks(block: dict[str, Any]) -> int:
    return 3 if _heading_level(block) >= 3 else 2


def _style_docx_heading(paragraph: Any, size: int, rgb: tuple[int, int, int],
                        underline: bool = False) -> None:
    from docx.shared import Pt, RGBColor
    paragraph.paragraph_format.space_before = Pt(14 if size >= 18 else 10)
    paragraph.paragraph_format.space_after = Pt(8)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(*rgb)
        run.underline = underline
        run.font.name = "Calibri"


def _style_docx_body(paragraph: Any, text: str) -> None:
    from docx.shared import Pt, RGBColor
    _add_runs(paragraph, text)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        if run.font.name != "Courier New":
            run.font.name = "Calibri"
        if run.font.size is None:
            run.font.size = Pt(11)
        if run.font.color.rgb is None:
            run.font.color.rgb = RGBColor(*_NAVY_RGB)


def _shade_cell(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _docx_cell(cell: Any, text: str, *, header: bool = False) -> None:
    from docx.shared import Pt, RGBColor
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11 if header else 10)
    run.bold = header
    if header:
        run.font.color.rgb = RGBColor(*_WHITE_RGB)
        _shade_cell(cell, _ACCENT_HEX)
    else:
        run.font.color.rgb = RGBColor(*_NAVY_RGB)


def _add_runs(paragraph: Any, text: str) -> None:
    value = text or ""
    pos = 0
    for match in _INLINE_RUN.finditer(value):
        if match.start() > pos:
            paragraph.add_run(value[pos:match.start()])
        if match.group(2) or match.group(3):
            run = paragraph.add_run(match.group(2) or match.group(3))
            run.bold = True
        elif match.group(4) or match.group(5):
            run = paragraph.add_run(match.group(4) or match.group(5))
            run.italic = True
        elif match.group(6):
            run = paragraph.add_run(match.group(6))
            run.font.name = "Courier New"
        else:
            paragraph.add_run(match.group(7) or "")
        pos = match.end()
    if pos < len(value):
        paragraph.add_run(value[pos:])
    if pos == 0 and not value:
        paragraph.add_run("")


def _pdf_table(pdf: Any, columns: list[str], rows: list[list[str]]) -> None:
    n = max(1, len(columns))
    data = [tuple(strip_markdown_inline(c) for c in columns)]
    for row in rows:
        data.append(tuple(strip_markdown_inline(c) for c in _pad(row, n)))
    try:
        from fpdf.fonts import FontFace
        heading = FontFace(emphasis="BOLD", color=_WHITE_RGB, fill_color=_ACCENT_RGB)
        with pdf.table(text_align="LEFT", headings_style=heading, first_row_as_headings=True) as table:
            for values in data:
                row = table.row()
                for value in values:
                    row.cell(_pdf_text(value))
        return
    except Exception:
        pass
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_text_color(*_WHITE_RGB)
    pdf.set_fill_color(*_ACCENT_RGB)
    pdf.multi_cell(0, 6, _pdf_text(" | ".join(data[0])), fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*_NAVY_RGB)
    for values in data[1:]:
        pdf.multi_cell(0, 6, _pdf_text(" | ".join(values)), new_x="LMARGIN", new_y="NEXT")


def log_ready(doc: dict[str, Any]) -> None:
    """Record that a document was accepted, without the body."""
    logger.info("artifact ready title_len=%s blocks=%s", len(doc["title"]), len(doc["blocks"]))
