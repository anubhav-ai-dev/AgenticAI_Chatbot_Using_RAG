"""Structured documents: validation, on-demand files, and the artifact SSE event."""

from __future__ import annotations

import json
from io import BytesIO

import pytest

from backend.services.artifact import (
    ArtifactError,
    content_disposition,
    document_from_markdown,
    document_source,
    is_bare_save,
    is_restyle,
    preferred_format,
    render_download,
    safe_filename,
    should_author_document,
    strip_markdown_inline,
    topic_title,
    validate_document,
    wants_file,
)
from backend.utils.rate_limiter import RateLimiter
from tests.conftest import FakeProviderClient, parse_sse

MODEL = "openai/gpt-oss-120b"


def new_session(tc, **kw):
    body = {"user_id": "u1", "title": "New Chat", "mode": "general", **kw}
    r = tc.post("/sessions", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def chat(tc, sid, message, **kw):
    payload = {"session_id": sid, "message": message, "user_id": "u1", "mode": "general",
               "web_search": True, **kw}
    r = tc.post("/chat/stream", json=payload)
    return r, parse_sse(r.text) if r.status_code == 200 else []


def names(events):
    return [e["event"] for e in events]


def text_of(events):
    return "".join(e["data"]["text"] for e in events if e["event"] == "token")


SAMPLE = {
    "title": "Café plan",
    "extra": "drop me",
    "blocks": [
        {"type": "heading", "text": "Overview\x00", "level": 9},
        {"type": "paragraph", "text": "Line one\nLine two", "html": "<script>alert(1)</script>"},
        {"type": "table", "columns": ["City", "Pop"], "rows": [["Pune", "1"]], "note": "nope"},
    ],
}


def test_validator_strips_controls_and_drops_unknown_fields():
    doc = validate_document(SAMPLE)
    assert set(doc) == {"title", "blocks"}
    assert doc["title"] == "Café plan"
    assert doc["blocks"][0] == {"type": "heading", "text": "Overview", "level": 2}
    assert "html" not in doc["blocks"][1]
    assert doc["blocks"][1]["text"] == "Line one\nLine two"
    assert doc["blocks"][2]["rows"] == [["Pune", "1"]]


def test_validator_rejects_empty_and_over_cap_documents():
    with pytest.raises(ArtifactError, match="empty"):
        validate_document({"title": "  ", "blocks": [{"type": "paragraph", "text": "x"}]})
    with pytest.raises(ArtifactError, match="empty"):
        validate_document({"title": "Notes", "blocks": []})
    with pytest.raises(ArtifactError, match="too long"):
        validate_document({"title": "x" * 121, "blocks": [{"type": "paragraph", "text": "ok"}]})
    with pytest.raises(ArtifactError, match="too many rows"):
        validate_document({
            "title": "Big",
            "blocks": [{"type": "table", "columns": ["A"], "rows": [["a"]] * 5001}],
        })
    with pytest.raises(ArtifactError, match="too long"):
        validate_document({
            "title": "Big",
            "blocks": [{"type": "paragraph", "text": "a" * 8001}],
        })


def test_when_a_document_is_and_is_not_authored():
    assert should_author_document("Explain recursion simply please", "A function that calls itself.") is False
    assert should_author_document("Write a document for the launch", "Here is the proposal.") is True
    assert should_author_document("Please write a report", "Which topic should it cover?") is False
    long = "# Overview\n\n" + ("Revenue grew. " * 40) + "\n\n# Outlook\n\nSteady next year."
    assert should_author_document("Tell me about the market in detail", long) is False
    assert wants_file("how to install omniroute in claude") is False
    assert wants_file("how to create a pdf") is True
    assert wants_file("what is markdown") is True
    assert wants_file("i want a pdf") is True
    assert is_bare_save("i want a pdf") is True
    assert preferred_format("i want a pdf") == "pdf"
    assert wants_file("generate pdf") is True
    assert wants_file("preview this as a document") is True
    assert wants_file("download as word") is True
    assert wants_file("write a proposal") is False
    assert wants_file("write a document") is True
    assert wants_file("please send the .xlxs") is True
    assert wants_file("text file") is True
    assert wants_file("password") is False
    assert wants_file("make the headings larger") is True
    assert is_restyle("make the headings larger") is True
    assert preferred_format("please send the .xlxs") == "xlsx"
    assert preferred_format("export as .pdf") == "pdf"
    assert preferred_format("save as a text file") == "txt"
    assert preferred_format("make a document") == "md"
    assert should_author_document("how to install X", long) is False
    assert should_author_document("preview this as a document", "Here is the document.") is True


def test_download_filename_uses_heading_with_underscores():
    assert safe_filename("Café plan", "md") == "Café_Plan.md"
    assert safe_filename("Launch proposal", "pdf") == "Launch_Proposal.pdf"
    assert safe_filename("AI Technology overview", "pdf") == "AI_Technology_Overview.pdf"
    assert safe_filename("The latest 10 AI trends this year", "pdf") == "Latest_10_AI_Trends.pdf"
    assert len(safe_filename("A very long heading about several extra topics", "docx").split("_")) <= 4
    assert " " not in safe_filename("Q3 market review", "docx")
    assert "-" not in safe_filename("AI-native development platforms", "pdf")


def test_five_formats_are_built_from_the_same_document():
    doc = validate_document(SAMPLE)
    md, md_type, md_name = render_download(doc, "md")
    txt, txt_type, txt_name = render_download(doc, "txt")
    xlsx, xlsx_type, xlsx_name = render_download(doc, "xlsx")
    docx, docx_type, docx_name = render_download(doc, "docx")
    pdf, pdf_type, pdf_name = render_download(doc, "pdf")

    assert md_name == safe_filename("Café plan", "md") and md_name.endswith(".md")
    assert "attachment;" in content_disposition(md_name)
    assert b"# Caf" in md and b"## Overview" in md and b"| City | Pop |" in md
    assert b"<script>" not in md
    assert txt_type.startswith("text/plain") and b"##" not in txt and b"Pune" in txt
    assert xlsx[:2] == b"PK" and xlsx_name.endswith(".xlsx")
    assert docx[:2] == b"PK" and docx_type.startswith("application/vnd.openxmlformats")
    assert pdf.startswith(b"%PDF") and pdf_type == "application/pdf" and pdf_name.endswith(".pdf")
    assert md_type.startswith("text/markdown")

    from openpyxl import load_workbook
    from docx import Document

    sheet = load_workbook(BytesIO(xlsx)).active
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert "Overview" in values and "Pune" in values and "Line one\nLine two" in values

    word = Document(BytesIO(docx))
    prose = "\n".join(p.text for p in word.paragraphs)
    assert "Café plan" in prose and "Overview" in prose
    assert word.tables[0].rows[1].cells[0].text == "Pune"


def test_ordinary_chat_emits_no_artifact(client_factory):
    tc, _ = client_factory()
    sid = new_session(tc)["id"]
    _, ev = chat(tc, sid, "Explain recursion simply please")
    assert names(ev)[0] == "meta" and names(ev)[-1] == "done"
    assert "artifact" not in names(ev)
    saved = tc.get(f"/sessions/{sid}").json()["messages"][-1]
    assert "artifact" not in (saved["meta"] or {})


ANSWER = (
    "# Launch proposal\n\n"
    "Ship the first version in March.\n\n"
    "| Item | Owner |\n"
    "| --- | --- |\n"
    "| API | Ada |\n\n"
    "```python\nprint(1)\n```\n"
)


def test_awkward_answers_still_become_every_format():
    one_col = document_from_markdown("| City |\n| --- |\n| Pune |\n", fallback_title="Cities")
    assert one_col["blocks"][0]["type"] == "table"
    assert one_col["blocks"][0]["columns"] == ["City"]
    assert one_col["blocks"][0]["rows"] == [["Pune"]]
    ruled = document_from_markdown("---\n", fallback_title="Notes")
    assert ruled["title"] == "Notes" and ruled["blocks"][0]["text"] == "---"
    wide = document_from_markdown(
        "| " + " | ".join(f"C{i}" for i in range(25)) + " |\n"
        "| " + " | ".join("---" for _ in range(25)) + " |\n"
        "| " + " | ".join("x" for _ in range(25)) + " |\n",
        fallback_title="Wide",
    )
    assert len(wide["blocks"][0]["columns"]) == 20
    for doc in (one_col, ruled, wide):
        for fmt in ("md", "txt", "xlsx", "docx", "pdf"):
            data, _, name = render_download(doc, fmt)
            assert data and name.endswith("." + fmt)


TRENDS = (
    "Most of what I found was marketing-grade listicles. The two sources I trust are Gartner and ABI.\n\n"
    "Multiagent systems. Gartner lists multiagent systems among its top trends for 2026. "
    "Think of a team of specialists. Caveat: ABI Research expects adoption to stay in pilot mode.\n\n"
    "AI-native development platforms. The development tool itself is built around AI.\n\n"
    "Physical AI. This is AI moving from screens into robots and machines.\n\n"
    "Evidence quality:\n\n"
    "Items 1 to 3 are analyst predictions, not measured outcomes.\n"
)


def test_a_prose_answer_becomes_a_titled_sectioned_document():
    assert topic_title("generate pdf", [
        {"role": "user", "content": "latest 10 ai trends"},
        {"role": "assistant", "content": TRENDS},
    ]) == "Latest 10 AI trends"
    doc = document_from_markdown(TRENDS, fallback_title="Latest 10 AI trends")
    assert doc["title"] == "Latest 10 AI trends"
    headings = [b["text"] for b in doc["blocks"] if b["type"] == "heading"]
    assert headings[:3] == [
        "1. Multiagent systems",
        "2. AI-native development platforms",
        "3. Physical AI",
    ]
    assert "Evidence quality" in headings
    caveat = next(b["text"] for b in doc["blocks"] if b["text"].startswith("Caveat:"))
    assert "pilot mode" in caveat
    for fmt in ("md", "txt", "xlsx", "docx", "pdf"):
        data, _, name = render_download(doc, fmt)
        assert data and name.endswith("." + fmt)
        if fmt == "pdf":
            assert data.startswith(b"%PDF")


def test_generate_pdf_saves_the_previous_answer(client_factory):
    groq = FakeProviderClient({MODEL: [
        ("ok", TRENDS),
        ("ok", "Your PDF is ready. It has the three trends, the ABI caveat, and the evidence note."),
    ]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    first, ev1 = chat(tc, sid, "latest 10 ai trends")
    assert first.status_code == 200 and "artifact" not in names(ev1)
    _, ev2 = chat(tc, sid, "generate pdf")
    assert "Do not repeat the full document" in groq.calls[-1]["messages"][0]["content"]
    artifact = next(e["data"] for e in ev2 if e["event"] == "artifact")
    assert artifact["format"] == "pdf"
    assert artifact["title"] == "Latest 10 AI trends"
    assert any(b.get("text") == "1. Multiagent systems" for b in artifact["blocks"])
    assert "Notepad" not in json.dumps(artifact)
    for fmt, marker in (("md", b"# Latest 10 AI trends"), ("txt", b"Latest 10 AI trends"),
                        ("xlsx", b"PK"), ("docx", b"PK"), ("pdf", b"%PDF")):
        resp = tc.get(f"/messages/{artifact['message_id']}/download", params={"format": fmt, "user_id": "u1"})
        assert resp.status_code == 200 and resp.content.startswith(marker), fmt
        if fmt in ("md", "txt"):
            assert b"Multiagent systems" in resp.content


def test_markdown_becomes_a_document_without_a_second_model_call():
    doc = document_from_markdown(ANSWER, fallback_title="unused")
    assert doc["title"] == "Launch proposal"
    kinds = [b["type"] for b in doc["blocks"]]
    assert kinds == ["paragraph", "table", "code"]
    assert doc["blocks"][1]["rows"] == [["API", "Ada"]]
    assert "print(1)" in doc["blocks"][2]["text"]
    previous = "# AI trends\n\n" + ("Models got faster. " * 20)
    how_to = "Open Notepad, then use Microsoft Print to PDF. Google Docs can export as PDF too."
    assert document_source("create a pdf", how_to, [
        {"role": "assistant", "content": previous},
    ]).startswith("# AI trends")
    assert document_source("Write a document for the launch", ANSWER, [
        {"role": "assistant", "content": previous},
    ]) == ANSWER.strip()
    refusal = (
        "I cannot directly generate or send a PDF file. However, I can provide you with "
        "the content formatted so you can easily copy it into a document and save it as a PDF."
    )
    assert document_source("i want a pdf", refusal, [
        {"role": "assistant", "content": previous},
    ]).startswith("# AI trends")


def test_file_request_streams_the_answer_then_the_artifact(client_factory):
    groq = FakeProviderClient({MODEL: [("ok", ANSWER)]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    _, ev = chat(tc, sid, "Write a document for the launch")
    assert names(ev)[0] == "meta" and names(ev)[-1] == "done"
    assert names(ev).index("artifact") < names(ev).index("done")
    assert "print(1)" in text_of(ev)
    assert "blocks" not in text_of(ev)
    artifact = next(e["data"] for e in ev if e["event"] == "artifact")
    assert artifact["title"] == "Launch proposal"
    assert artifact["format"] == "md"
    assert len(groq.calls) == 1
    assert "downloadable file" in groq.calls[0]["messages"][0]["content"]

    saved = tc.get(f"/sessions/{sid}").json()["messages"][-1]
    assert saved["id"] == artifact["message_id"]
    assert saved["meta"]["artifact"]["blocks"][1]["rows"] == [["API", "Ada"]]
    assert "PK" not in json.dumps(saved["meta"])

    resp = tc.get(f"/messages/{saved['id']}/download", params={"format": "md", "user_id": "u1"})
    assert resp.status_code == 200
    assert resp.headers["content-disposition"].startswith("attachment;")
    assert "Launch_Proposal.md" in resp.headers["content-disposition"]
    assert b"# Launch proposal" in resp.content and b"print(1)" in resp.content
    assert tc.get(f"/messages/{saved['id']}/download", params={"format": "pdf", "user_id": "other"}).status_code == 404
    plain = tc.get(f"/messages/{saved['id']}/download", params={"format": "txt", "user_id": "u1"})
    sheet = tc.get(f"/messages/{saved['id']}/download", params={"format": "xlsx", "user_id": "u1"})
    word = tc.get(f"/messages/{saved['id']}/download", params={"format": "docx", "user_id": "u1"})
    pdf = tc.get(f"/messages/{saved['id']}/download", params={"format": "pdf", "user_id": "u1"})
    assert plain.content.startswith(b"Launch proposal") and b"##" not in plain.content and b"print(1)" in plain.content
    assert sheet.content[:2] == b"PK" and word.content[:2] == b"PK" and pdf.content.startswith(b"%PDF")


def test_save_request_files_the_previous_answer_not_export_instructions(client_factory):
    trends = "# AI trends\n\n" + ("Frontier models got faster and cheaper. " * 12)
    how_to = "Copy the text into Notepad, then choose Microsoft Print to PDF."
    groq = FakeProviderClient({MODEL: [("ok", trends), ("ok", how_to)]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    first, ev1 = chat(tc, sid, "What are the main AI trends this year?")
    assert first.status_code == 200 and "artifact" not in names(ev1)
    _, ev2 = chat(tc, sid, "create a pdf")
    artifact = next(e["data"] for e in ev2 if e["event"] == "artifact")
    assert artifact["title"] == "AI trends"
    assert artifact["format"] == "pdf"
    assert "Frontier models" in artifact["blocks"][0]["text"]
    assert "Notepad" not in json.dumps(artifact)
    pdf = tc.get(f"/messages/{artifact['message_id']}/download", params={"format": "pdf", "user_id": "u1"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_i_want_a_pdf_opens_a_file_even_when_the_model_refuses(client_factory):
    body = "# 2026 Calendar\n\n" + ("Printable templates are listed by month. " * 12)
    refusal = (
        "I cannot directly generate or send a PDF file. However, I can provide you with "
        "the content formatted so you can easily copy it into a document and save it as a PDF."
    )
    groq = FakeProviderClient({MODEL: [("ok", body), ("ok", refusal)]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    first, ev1 = chat(tc, sid, "What calendars are available?")
    assert first.status_code == 200 and "artifact" not in names(ev1)
    _, ev2 = chat(tc, sid, "i want a pdf")
    artifact = next(e["data"] for e in ev2 if e["event"] == "artifact")
    assert artifact["format"] == "pdf"
    assert artifact["title"] == "2026 Calendar"
    assert "Printable templates" in json.dumps(artifact)
    assert "cannot directly generate" not in json.dumps(artifact).lower()


@pytest.mark.parametrize("mode", ["general", "document", "research", "code"])
def test_every_mode_downloads_every_format(client_factory, mode):
    groq = FakeProviderClient({MODEL: [("ok", ANSWER)]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc, mode=mode)["id"]
    _, ev = chat(tc, sid, "Create a pdf of these notes", mode=mode)
    assert names(ev)[0] == "meta" and names(ev)[-1] == "done"
    assert ev[0]["data"]["mode"] == mode
    artifact = next(e["data"] for e in ev if e["event"] == "artifact")
    assert artifact["title"] == "Launch proposal"
    assert artifact["format"] == "pdf"
    assert "downloadable file" in groq.calls[0]["messages"][0]["content"]
    saved_id = artifact["message_id"]
    files = {}
    for fmt in ("md", "txt", "xlsx", "docx", "pdf"):
        resp = tc.get(f"/messages/{saved_id}/download", params={"format": fmt, "user_id": "u1"})
        assert resp.status_code == 200, (mode, fmt, resp.text)
        assert "attachment;" in resp.headers["content-disposition"]
        files[fmt] = resp.content
    assert files["md"].startswith(b"# Launch proposal") and b"print(1)" in files["md"]
    assert files["txt"].startswith(b"Launch proposal") and b"Ship the first version" in files["txt"]
    assert b"##" not in files["txt"] and b"print(1)" in files["txt"]
    assert files["xlsx"].startswith(b"PK") and files["docx"].startswith(b"PK")
    assert files["pdf"].startswith(b"%PDF")
    from openpyxl import load_workbook
    from docx import Document
    sheet = [c.value for row in load_workbook(BytesIO(files["xlsx"])).active.iter_rows() for c in row if c.value]
    assert "API" in sheet and "Ada" in sheet and "print(1)" in "\n".join(map(str, sheet))
    prose = "\n".join(p.text for p in Document(BytesIO(files["docx"])).paragraphs)
    assert "Launch proposal" in prose and "print(1)" in prose


def test_thin_answer_still_downloads_instead_of_failing(client_factory):
    groq = FakeProviderClient({MODEL: [("ok", "---\n")]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    status, ev = chat(tc, sid, "create a pdf")
    assert status.status_code == 200 and names(ev)[-1] == "done"
    artifact = next(e["data"] for e in ev if e["event"] == "artifact")
    assert "error" not in artifact and artifact["format"] == "pdf"
    for fmt in ("md", "txt", "xlsx", "docx", "pdf"):
        resp = tc.get(f"/messages/{artifact['message_id']}/download", params={"format": fmt, "user_id": "u1"})
        assert resp.status_code == 200, (fmt, resp.text)


def test_clarification_does_not_create_a_file(client_factory):
    groq = FakeProviderClient({MODEL: [("ok", "Which topic should the report cover?")]})
    tc, _ = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    _, ev = chat(tc, sid, "Please write a report")
    assert "artifact" not in names(ev)
    assert len(groq.calls) == 1


def test_download_rate_limit_is_separate_from_chat(client_factory):
    groq = FakeProviderClient({MODEL: [("ok", "Notes are ready.")]})
    tc, c = client_factory(llm_client=groq)
    sid = new_session(tc)["id"]
    _, ev = chat(tc, sid, "Make notes for the meeting as a document")
    mid = ev[-1]["data"]["message_id"]
    c.download_limiter = RateLimiter(1)
    assert tc.get(f"/messages/{mid}/download", params={"format": "txt", "user_id": "u1"}).status_code == 200
    limited = tc.get(f"/messages/{mid}/download", params={"format": "txt", "user_id": "u1"})
    assert limited.status_code == 429 and int(limited.headers["retry-after"]) >= 1
    # The chat bucket was not the one that ran out.
    assert tc.post("/chat/stream", json={
        "session_id": sid, "message": "Explain recursion simply please", "user_id": "u1",
    }).status_code == 200


def test_inline_markdown_is_stripped_in_non_markdown_formats():
    doc = document_from_markdown(
        "# Guide\n\nUse **Create Key** and *optional* `token`.\n\n"
        "- First item\n- Second item\n",
        fallback_title="Guide",
    )
    assert any(b["type"] == "list" for b in doc["blocks"])
    md, _, _ = render_download(doc, "md")
    assert b"**Create Key**" in md
    txt, _, _ = render_download(doc, "txt")
    assert b"**" not in txt and b"Create Key" in txt and b"##" not in txt
    xlsx, _, _ = render_download(doc, "xlsx")
    from openpyxl import load_workbook
    values = [c.value for row in load_workbook(BytesIO(xlsx)).active.iter_rows() for c in row if c.value]
    joined = "\n".join(map(str, values))
    assert "**" not in joined and "Create Key" in joined
    word_bytes, _, _ = render_download(doc, "docx")
    from docx import Document
    prose = "\n".join(p.text for p in Document(BytesIO(word_bytes)).paragraphs)
    assert "**" not in prose and "Create Key" in prose
    pdf, _, _ = render_download(doc, "pdf")
    assert pdf.startswith(b"%PDF")
    assert strip_markdown_inline("Use **Create Key** and [site](https://x.y)") == "Use Create Key and site"
