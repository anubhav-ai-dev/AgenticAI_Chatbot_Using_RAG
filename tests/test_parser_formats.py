"""Parser coverage for the document/image formats beyond the original PDF/DOCX/CSV/XLSX/TXT set."""

from __future__ import annotations

import json
import zipfile

import pytest

from backend.core.exceptions import DocumentParseError
from backend.services import parser as parser_mod
from backend.services.embedding_service import EmbeddingService
from backend.services.parser import DocumentParser, SUPPORTED_EXTENSIONS, ocr_available

PHRASE = "Refund policy thirty days"


def _text(pages) -> str:
    return "\n".join(p["text"] for p in pages)


@pytest.fixture
def parser():
    return DocumentParser()


def _image(tmp_path, name: str, fmt: str):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (900, 140), "white")
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 44)
    except OSError:
        font = ImageFont.load_default()
    ImageDraw.Draw(img).text((20, 40), PHRASE, fill="black", font=font)
    path = tmp_path / name
    img.save(path, fmt)
    return path


def test_supported_extensions_cover_images_and_office_formats():
    assert {"jpg", "jpeg", "png", "pptx", "html", "json", "rtf", "odt", "xls"} <= SUPPORTED_EXTENSIONS


def test_unsupported_extension_still_rejected(parser, tmp_path):
    f = tmp_path / "a.exe"
    f.write_bytes(b"MZ")
    with pytest.raises(DocumentParseError, match="Unsupported file type"):
        parser.parse_file(str(f), "a.exe")


def test_pptx_slides_notes_and_tables(parser, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Quarterly results"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Q", "Rev", "Q1", "42"
    s.notes_slide.notes_text_frame.text = "speaker note here"
    path = tmp_path / "d.pptx"
    prs.save(path)
    pages = parser.parse_file(str(path), "d.pptx")
    assert pages[0]["page_number"] == 1
    t = _text(pages)
    assert "Quarterly results" in t and "Q1\t42" in t and "speaker note here" in t


def test_html_file_extracts_text_without_scripts(parser, tmp_path):
    f = tmp_path / "p.html"
    f.write_text("<html><body><script>evil()</script><h1>Title</h1><p>" + "Body text. " * 30
                 + "</p></body></html>", encoding="utf-8")
    t = _text(parser.parse_file(str(f), "p.html"))
    assert "Body text" in t and "evil()" not in t


def test_json_is_flattened_to_text(parser, tmp_path):
    f = tmp_path / "d.json"
    f.write_text(json.dumps({"policy": {"refund_days": 30}, "items": ["a", "b"]}), encoding="utf-8")
    t = _text(parser.parse_file(str(f), "d.json"))
    assert '"refund_days": 30' in t


def test_invalid_json_is_a_parse_error(parser, tmp_path):
    f = tmp_path / "bad.json"
    f.write_text("{not json", encoding="utf-8")
    with pytest.raises(DocumentParseError):
        parser.parse_file(str(f), "bad.json")


def test_rtf(parser, tmp_path):
    f = tmp_path / "d.rtf"
    f.write_bytes(rb"{\rtf1\ansi Refund policy thirty days\par Second paragraph.}")
    assert "Refund policy thirty days" in _text(parser.parse_file(str(f), "d.rtf"))


def test_odt(parser, tmp_path):
    f = tmp_path / "d.odt"
    xml = ('<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
           '<text:h>Heading</text:h><text:p>Refund policy thirty days</text:p>'
           '</office:text></office:body></office:document-content>')
    with zipfile.ZipFile(f, "w") as zf:
        zf.writestr("content.xml", xml)
    t = _text(parser.parse_file(str(f), "d.odt"))
    assert "Heading" in t and "Refund policy thirty days" in t


def test_xls_legacy_excel(parser, tmp_path):
    xlwt = pytest.importorskip("xlwt")                  # only to build a fixture
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    ws.write(0, 0, "name")
    ws.write(1, 0, "alice")
    path = tmp_path / "d.xls"
    wb.save(str(path))
    assert "alice" in _text(parser.parse_file(str(path), "d.xls"))


def test_non_utf8_text_is_decoded_not_mojibake(parser, tmp_path):
    f = tmp_path / "t.txt"
    f.write_bytes("Café résumé naïve — prix 5€ ".encode("cp1252") * 5)
    t = _text(parser.parse_file(str(f), "t.txt"))
    assert "Café" in t and "�" not in t


def test_utf16_text(parser, tmp_path):
    f = tmp_path / "t.txt"
    f.write_bytes(("Refund policy thirty days\n\nSecond section here. " * 3).encode("utf-16"))
    assert "Refund policy thirty days" in _text(parser.parse_file(str(f), "t.txt"))


needs_ocr = pytest.mark.skipif(not ocr_available(), reason="OCR engine not installed")


@needs_ocr
@pytest.mark.parametrize("name,fmt", [("a.png", "PNG"), ("a.jpg", "JPEG"), ("a.jpeg", "JPEG"),
                                      ("a.webp", "WEBP"), ("a.bmp", "BMP")])
def test_image_text_is_extracted_by_ocr(parser, tmp_path, name, fmt):
    path = _image(tmp_path, name, fmt)
    pages = parser.parse_file(str(path), name)
    assert "refund" in _text(pages).lower() and pages[0]["metadata"]["ocr"] is True


@needs_ocr
def test_image_without_text_is_a_clear_error(parser, tmp_path):
    from PIL import Image

    path = tmp_path / "blank.png"
    Image.new("RGB", (300, 300), "white").save(path)
    with pytest.raises(DocumentParseError, match="No readable text"):
        parser.parse_file(str(path), "blank.png")


@needs_ocr
def test_scanned_pdf_pages_fall_back_to_ocr(parser, tmp_path):
    img = _image(tmp_path, "scan.png", "PNG")
    from PIL import Image

    pdf = tmp_path / "scan.pdf"
    Image.open(img).convert("RGB").save(pdf, "PDF")           # image-only PDF, no text layer
    pages = parser.parse_file(str(pdf), "scan.pdf")
    assert pages and "refund" in _text(pages).lower()
    assert pages[0]["metadata"].get("ocr") is True


def test_image_upload_reports_unavailable_ocr(parser, tmp_path, monkeypatch):
    monkeypatch.setattr(parser_mod, "_get_ocr", lambda: None)
    path = tmp_path / "a.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(DocumentParseError, match="OCR"):
        parser.parse_file(str(path), "a.png")


def test_embedding_load_failure_is_retried_after_cooldown(settings, monkeypatch):
    from tests.conftest import FakeEmbedder

    calls = {"n": 0}

    def factory(_s):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("offline")
        return FakeEmbedder()

    svc = EmbeddingService(settings, client_factory=factory)
    assert svc._load() is None and svc.error and calls["n"] == 1
    assert svc._load() is None and calls["n"] == 1                 # still cooling down
    svc._failed_at -= 120                                          # cooldown elapsed
    assert svc._load() is not None and calls["n"] == 2 and svc.error is None and svc.ready


@needs_ocr
def test_png_upload_is_indexed_and_retrievable(client_factory, tmp_path):
    path = _image(tmp_path, "receipt.png", "PNG")
    tc, _ = client_factory()
    r = tc.post("/upload?user_id=u1", files={"file": ("receipt.png", path.read_bytes(), "image/png")})
    assert r.status_code == 200, r.text
    assert r.json()["file_type"] == "png" and r.json()["chunk_count"] >= 1


def _text_pdf(pages: int) -> bytes:
    kids, body = [], []
    n = 3
    for i in range(pages):
        stream = f"BT /F1 12 Tf 40 700 Td (Refund policy page {i + 1} covers thirty days of returns) Tj ET"
        body.append((n, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {n + 1} 0 R "
                        "/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>"))
        body.append((n + 1, f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"))
        kids.append(f"{n} 0 R")
        n += 2
    objs = [(1, "<< /Type /Catalog /Pages 2 0 R >>"),
            (2, f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {pages} >>")] + body
    out, offs = b"%PDF-1.4\n", {}
    for num, txt in objs:
        offs[num] = len(out)
        out += f"{num} 0 obj\n{txt}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {n}\n0000000000 65535 f \n".encode()
    for num in range(1, n):
        out += f"{offs[num]:010d} 00000 n \n".encode()
    return out + f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


def test_text_pdf_uses_fast_pdfium_path_and_keeps_page_numbers(parser, tmp_path, monkeypatch):
    path = tmp_path / "t.pdf"
    path.write_bytes(_text_pdf(3))
    monkeypatch.setattr(DocumentParser, "_pdf_pdfplumber", staticmethod(lambda p: pytest.fail("slow parser used")))
    pages = parser.parse_file(str(path), "t.pdf")
    assert [p["page_number"] for p in pages] == [1, 2, 3]
    assert pages[0]["metadata"]["parser"] == "pdfium" and "Refund policy page 2" in pages[1]["text"]


def test_pdf_falls_back_to_pdfplumber_when_pdfium_fails(parser, tmp_path, monkeypatch):
    path = tmp_path / "t.pdf"
    path.write_bytes(_text_pdf(1))

    def boom(_p):
        raise RuntimeError("pdfium broke")

    monkeypatch.setattr(DocumentParser, "_pdf_pdfium", staticmethod(boom))
    pages = parser.parse_file(str(path), "t.pdf")
    assert pages and pages[0]["metadata"]["parser"] == "pdfplumber"


def test_scanned_pdf_is_not_reparsed_by_other_text_parsers(parser, tmp_path, monkeypatch):
    from PIL import Image

    path = tmp_path / "scan.pdf"
    Image.new("RGB", (200, 200), "white").save(path, "PDF")
    calls = []
    monkeypatch.setattr(DocumentParser, "_pdf_pdfplumber", staticmethod(lambda p: calls.append("plumber")))
    monkeypatch.setattr(DocumentParser, "_pdf_pypdf2", staticmethod(lambda p: calls.append("pypdf2")))
    monkeypatch.setattr(parser_mod, "_get_ocr", lambda: None)
    assert parser.parse_file(str(path), "scan.pdf") == [] and calls == []


def test_garbage_pdf_is_a_parse_error(parser, tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(DocumentParseError, match="could not be read"):
        parser.parse_file(str(path), "bad.pdf")


def test_embeddings_keep_input_order_after_length_sorting():
    import numpy as np
    from backend.client.embedding_client import LocalEmbeddingClient

    class StubModel:
        def embed(self, texts, batch_size=64):
            for t in texts:                         # vector encodes the text length
                yield np.array([float(len(t)), 1.0], dtype=np.float32)

    client = LocalEmbeddingClient.__new__(LocalEmbeddingClient)
    client._model, client._batch = StubModel(), 8
    texts = ["longer text here", "a", "mid text"]
    out = client.embed_documents(texts)
    # normalised (len, 1): x grows with length, so row order must match the input order
    assert out[0, 0] > out[2, 0] > out[1, 0]
    assert client.embed_documents([]).shape[0] == 0


def test_upload_logs_stage_timings(client_factory, caplog):
    import logging

    tc, _ = client_factory()
    with caplog.at_level(logging.INFO, logger="backend.services.document_service"):
        r = tc.post("/upload?user_id=u1", files={"file": ("a.txt", b"Refund policy thirty days. " * 20, "text/plain")})
    assert r.status_code == 200
    assert any("Upload timings txt" in m and "parse=" in m and "embed=" in m for m in caplog.messages)
